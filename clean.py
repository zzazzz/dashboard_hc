"""
HC Analytics Case 1 - Data Cleaning Pipeline
============================================
Input : HC_Analytics_Case1_Data_Dummy.xlsx (7 sheet)
Output: HC_Analytics_Case1_Data_Clean.xlsx (7 sheet bersih + 1 sheet Log Pembersihan)

Prinsip yang dipakai:
- Masalah format/typo  -> DIPERBAIKI
- Masalah yang butuh keputusan bisnis -> DITANDAI pakai kolom flag_*, tidak diubah
- Tidak ada baris yang dihapus kecuali benar-benar duplikat identik atau kosong total

Dependency: pandas, numpy, openpyxl
"""

import re
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC = "HC_Analytics_Case1_Data_Dummy.xlsx"
DST = "HC_Analytics_Case1_Data_Clean.xlsx"

# ---------------------------------------------------------------- referensi
# Nilai kanonik: sumber kebenaran untuk standardisasi.
# Semua varian kotor dipetakan ke sini lewat kunci lowercase.
REGIONAL_KANONIK = [
    "Regional 1 - Jakarta",
    "Regional 2 - Jawa Barat",
    "Regional 3 - Jawa Tengah & Jawa Timur",
    "Regional 4 - Sumatera",
    "Regional 5 - Wilayah Timur",
]

JOB_KANONIK = [
    "Marketing Officer", "Senior Marketing Officer", "Marketing Supervisor", "Sales Section Head",
    "Collection Officer", "Senior Collection Officer", "Collection Supervisor", "Collection Section Head",
    "Credit Analyst", "Senior Credit Analyst", "Credit Supervisor", "Credit Section Head",
    "Admin Officer", "Senior Admin Officer", "Admin Supervisor", "Operation Section Head",
]

ALASAN_KANONIK = [
    "Mendapat pekerjaan lain", "Beban kerja terlalu berat", "Tidak ada kejelasan jenjang karir",
    "Hubungan dengan atasan", "Kompensasi kurang kompetitif", "Pindah domisili", "Alasan keluarga",
    "Wirausaha", "Melanjutkan studi", "Tidak lulus masa percobaan", "Kinerja di bawah standar",
    "Pelanggaran disiplin",
]

# nilai_performa: 4 nilai salah format desimal (koma geser 1 digit)
KOREKSI_PERFORMA = {33.1: 3.31, 22.9: 2.29, 35.0: 3.50, 22.1: 2.21}

# Band kategori_performa.
# Konvensi: batas atas INKLUSIF -> naik band kalau nilai MELEBIHI cutoff (x > cutoff).
# Dasarnya bukan selera, tapi satu-satunya batas yang konsisten di data sumber:
#   nilai 4.20 -> selalu "Baik" (6 baris, 0 "Istimewa")
#   nilai 4.21 -> selalu "Istimewa" (3 baris, 0 "Baik")
# Batas 2.60 dan 3.40 di data sumber terbelah dua label, jadi konvensi dari
# batas 4.20 di atas yang diterapkan seragam ke semua batas.
BAND_PERFORMA = [
    (2.60, "Kurang"),    # 1.00 <= x <= 2.60
    (3.40, "Cukup"),     # 2.60 <  x <= 3.40
    (4.20, "Baik"),      # 3.40 <  x <= 4.20
    (float("inf"), "Istimewa"),  # x > 4.20
]


# ---------------------------------------------------------------- utilitas
def norm(s):
    """Rapikan string: buang spasi awal/akhir, rapatkan spasi ganda jadi satu."""
    if pd.isna(s):
        return np.nan
    return re.sub(r"\s+", " ", str(s)).strip()


def buat_mapper(nilai_kanonik, hapus_titik_akhir=False):
    """Bikin dict {bentuk_ternormalisasi_lowercase: nilai_kanonik}."""
    return {norm(v).lower(): v for v in nilai_kanonik}


def ke_kanonik(seri, mapper, hapus_titik_akhir=False):
    """Petakan seri kotor ke nilai kanonik. Gagal petakan -> NaN (sengaja, biar ketahuan)."""
    kunci = seri.map(norm)
    kunci = kunci.map(lambda v: np.nan if pd.isna(v) else v.rstrip(".").lower()
                      if hapus_titik_akhir else str(v).lower())
    hasil = kunci.map(mapper)
    # Safety net: kalau ada varian kotor yang belum terdaftar, script berhenti
    belum_terpetakan = sorted(set(kunci.dropna()) - set(mapper))
    assert not belum_terpetakan, f"Varian belum terdaftar di daftar kanonik: {belum_terpetakan}"
    return hasil


def kategorikan_performa(nilai):
    """Petakan nilai_performa (skala 1-5) ke kategori sesuai BAND_PERFORMA."""
    if pd.isna(nilai):
        return np.nan
    for cutoff, label in BAND_PERFORMA:
        if nilai <= cutoff:
            return label
    return BAND_PERFORMA[-1][1]


class LogPembersihan:
    """Pencatat setiap tindakan pembersihan, dikeluarkan jadi sheet dokumentasi."""

    def __init__(self):
        self.baris = []

    def catat(self, sheet, kolom, masalah, jumlah, aksi):
        self.baris.append({
            "sheet": sheet, "kolom": kolom, "masalah": masalah,
            "jumlah_baris": int(jumlah), "aksi": aksi,
        })

    def to_frame(self):
        return pd.DataFrame(self.baris)


# ---------------------------------------------------------------- 1. Data Cabang
def bersihkan_data_cabang(df, log):
    df = df.copy()
    for c in ["Branch Code", "Branch Name", "Kota"]:
        df[c] = df[c].map(norm)

    sebelum = df["Regional"].map(norm)
    df["Regional"] = ke_kanonik(df["Regional"], buat_mapper(REGIONAL_KANONIK))
    n = int((sebelum != df["Regional"]).sum())
    log.catat("Data Cabang", "Regional",
              'Inkonsistensi huruf besar/kecil (mis. "REGIONAL 2 - JAWA BARAT")',
              n, "Distandarkan ke 5 nilai kanonik")
    return df


# ---------------------------------------------------------------- 2. Data Karyawan
def bersihkan_data_karyawan(df, df_kompensasi, log):
    df = df.copy()

    # 2a. Baris kosong total (semua kolom NaN)
    n = int(df["NIK"].isna().sum())
    df = df[df["NIK"].notna()].copy()
    log.catat("Data Karyawan", "(semua)", "Baris kosong total", n, "Dihapus")

    # 2b. NIK duplikat — dicek dulu benar-benar identik atau tidak
    dup = df[df.duplicated("NIK", keep=False)]
    assert dup.drop_duplicates().shape[0] == dup["NIK"].nunique(), \
        "Ada NIK duplikat dengan isi BERBEDA - jangan auto-drop, cek manual!"
    n = int(df.duplicated("NIK", keep="first").sum())
    df = df.drop_duplicates("NIK", keep="first").copy()
    log.catat("Data Karyawan", "NIK", "NIK duplikat (baris identik persis)",
              n, "Dihapus, disisakan 1 baris per NIK")

    # 2c. Tanggal Masuk: tipe campuran (datetime + string "DD-Mon-YYYY")
    n = int(df["Tanggal Masuk"].map(lambda v: isinstance(v, str)).sum())
    df["Tanggal Masuk"] = pd.to_datetime(df["Tanggal Masuk"], errors="coerce",
                                         format="mixed", dayfirst=True)
    assert df["Tanggal Masuk"].isna().sum() == 0, "Ada tanggal masuk gagal di-parse"
    log.catat("Data Karyawan", "Tanggal Masuk",
              'Format campuran: teks "DD-Mon-YYYY" di tengah kolom datetime',
              n, "Di-parse ke datetime64 (YYYY-MM-DD)")

    # 2d. Tanggal Lahir placeholder 1900-01-01 (Excel serial 1)
    mask = df["Tanggal Lahir"].dt.year == 1900
    n = int(mask.sum())
    df.loc[mask, "Tanggal Lahir"] = pd.NaT
    log.catat("Data Karyawan", "Tanggal Lahir", "Placeholder 1900-01-01 (Excel serial 1)",
              n, "Diubah jadi NULL (tidak ada dasar imputasi)")

    # 2e. Job: 56 varian -> 16 kanonik
    sebelum = df["Job"].copy()
    df["Job"] = ke_kanonik(df["Job"], buat_mapper(JOB_KANONIK))
    n = int((sebelum != df["Job"]).sum())
    log.catat("Data Karyawan", "Job", "Spasi ganda, spasi di awal, UPPERCASE, lowercase",
              n, f"Distandarkan ke {len(JOB_KANONIK)} nama jabatan kanonik")

    # 2f. Branch Code kosong -> imputasi dari Data Kompensasi
    #     Aman karena tiap NIK cuma punya 1 kode cabang di sana (dicek di bawah)
    cek = df_kompensasi.groupby("nik")["kode_cabang"].nunique()
    peta = df_kompensasi.groupby("nik")["kode_cabang"].first()
    mask = df["Branch Code"].isna()
    nik_kosong = df.loc[mask, "NIK"]
    assert (cek.reindex(nik_kosong).fillna(0) == 1).all(), \
        "Ada NIK dengan >1 kode cabang di Kompensasi - imputasi tidak aman"
    n = int(mask.sum())
    df.loc[mask, "Branch Code"] = nik_kosong.map(peta)
    log.catat("Data Karyawan", "Branch Code", "Kosong", n,
              "Diimputasi dari Data Kompensasi (kode cabang unik per NIK)")

    for c in ["NIK", "Branch Code", "Function", "Level/Pangkat",
              "Jenis Kelamin", "Pendidikan", "Status Kekaryawan"]:
        df[c] = df[c].map(norm)

    # 2g. FLAG: umur saat masuk tidak wajar (tidak diubah, butuh konfirmasi HC)
    umur = (df["Tanggal Masuk"] - df["Tanggal Lahir"]).dt.days / 365.25
    df["flag_umur_masuk_tidak_wajar"] = (umur < 18).fillna(False)
    log.catat("Data Karyawan", "Tanggal Lahir vs Tanggal Masuk",
              "Umur saat masuk < 18 tahun (min 8,3 th)",
              int(df["flag_umur_masuk_tidak_wajar"].sum()),
              "DITANDAI saja (flag), tidak diubah")

    return df


# ---------------------------------------------------------------- 3. Reason Out
def bersihkan_reason_out(df, nik_master, log):
    df = df.copy()

    n = int(df.duplicated(keep="first").sum())
    df = df.drop_duplicates(keep="first").copy()
    log.catat("Reason Out", "(semua)", "Baris duplikat identik", n, "Dihapus")

    # Alasan keluar: 46 varian -> 12 kanonik (titik di akhir, spasi awal, casing campur)
    n_null = int(df["alasan_keluar_exit_interview"].isna().sum())
    sebelum = df["alasan_keluar_exit_interview"].copy()
    bersih = ke_kanonik(df["alasan_keluar_exit_interview"],
                        buat_mapper(ALASAN_KANONIK), hapus_titik_akhir=True)
    n_ubah = int((bersih.fillna("~") != sebelum.fillna("~")).sum()) - n_null
    df["alasan_keluar_exit_interview"] = bersih.fillna("Tidak Diisi")
    log.catat("Reason Out", "alasan_keluar_exit_interview",
              "Titik di akhir, spasi awal, UPPERCASE, lowercase",
              n_ubah, f"Distandarkan ke {len(ALASAN_KANONIK)} alasan kanonik")
    log.catat("Reason Out", "alasan_keluar_exit_interview", "Nilai kosong",
              n_null, 'Diisi "Tidak Diisi"')

    for c in ["nik", "kode_cabang", "fungsi", "level", "kategori_keluar"]:
        df[c] = df[c].map(norm)
    df["jabatan"] = ke_kanonik(df["jabatan"], buat_mapper(JOB_KANONIK))

    # Kategori performa terakhir diturunkan pakai band yang SAMA dengan
    # Performance Data, supaya taksonomi antar sheet konsisten saat di-join.
    df["kategori_performa_terakhir"] = df["nilai_performa_terakhir"].map(kategorikan_performa)
    log.catat("Reason Out", "nilai_performa_terakhir",
              "Tidak ada kolom kategori (sheet lain punya)", len(df),
              "Ditambah kolom turunan kategori_performa_terakhir pakai band yang sama")

    # FLAG: NIK yatim (ada di Reason Out tapi tidak ada di master karyawan)
    df["flag_nik_tidak_ada_di_master"] = ~df["nik"].isin(nik_master)
    log.catat("Reason Out", "nik", "NIK tidak ada di Data Karyawan (E999001-E999004)",
              int(df["flag_nik_tidak_ada_di_master"].sum()),
              "DITANDAI saja (flag), tidak dihapus")

    return df


# ---------------------------------------------------------------- 4. Performance Data
def bersihkan_performance_data(df, log):
    df = df.copy()

    n = int(df.duplicated(keep="first").sum())
    df = df.drop_duplicates(keep="first").copy()
    log.catat("Performance Data", "(semua)", "Baris duplikat identik", n, "Dihapus")

    # Skala seharusnya 1-5; nilai >5 = salah format desimal
    mask = df["nilai_performa"] > 5
    tak_dikenal = set(df.loc[mask, "nilai_performa"]) - set(KOREKSI_PERFORMA)
    assert not tak_dikenal, f"Nilai >5 di luar daftar koreksi: {tak_dikenal}"
    n = int(mask.sum())
    df.loc[mask, "nilai_performa"] = df.loc[mask, "nilai_performa"].map(KOREKSI_PERFORMA)
    assert df["nilai_performa"].between(1, 5).all(), "Masih ada nilai di luar skala 1-5"
    log.catat("Performance Data", "nilai_performa",
              "Salah format desimal, nilai > 5 (33.1 / 22.9 / 35.0 / 22.1)",
              n, "Dikoreksi ke 3.31 / 2.29 / 3.50 / 2.21")

    # kategori_performa dihitung ULANG dari nilai_performa pakai BAND_PERFORMA,
    # bukan dipercaya apa adanya. Label lama disimpan buat audit trail.
    df["kategori_performa_sumber"] = df["kategori_performa"]
    df["kategori_performa"] = df["nilai_performa"].map(kategorikan_performa)
    n = int((df["kategori_performa"] != df["kategori_performa_sumber"]).sum())
    log.catat("Performance Data", "kategori_performa",
              "Label tidak konsisten di batas band (2.60 -> Kurang/Cukup; 3.40 -> Cukup/Baik)",
              n, "Dihitung ulang dari nilai_performa pakai band tunggal "
                 "(<=2.60 Kurang, <=3.40 Cukup, <=4.20 Baik, >4.20 Istimewa); "
                 "label asli disimpan di kolom kategori_performa_sumber")

    return df


# ---------------------------------------------------------------- 5. Data Kompensasi
def bersihkan_data_kompensasi(df, log):
    df = df.copy()

    # Gaji negatif: cek dulu apakah |nilai| cocok dengan bulan lain NIK yang sama.
    # Kalau cocok -> murni salah tanda, aman di-abs().
    mask = df["gaji_pokok_juta"] < 0
    for nik, nilai in df.loc[mask, ["nik", "gaji_pokok_juta"]].itertuples(index=False):
        lain = df[(df["nik"] == nik) & (df["gaji_pokok_juta"] > 0)]["gaji_pokok_juta"]
        assert abs(nilai) in set(lain), f"{nik}: nilai negatif tidak cocok riwayat gaji"
    n = int(mask.sum())
    df.loc[mask, "gaji_pokok_juta"] = df.loc[mask, "gaji_pokok_juta"].abs()
    log.catat("Data Kompensasi", "gaji_pokok_juta",
              "Nilai negatif (salah tanda; bulan lain NIK sama bernilai positif sama)",
              n, "Diubah ke nilai absolut")

    # insentif_realisasi kosong: JANGAN diisi 0 - "tidak dapat insentif" beda arti
    # dengan "data tidak tercatat". Cukup ditandai.
    df["flag_insentif_realisasi_kosong"] = df["insentif_realisasi_juta"].isna()
    log.catat("Data Kompensasi", "insentif_realisasi_juta", "Nilai kosong",
              int(df["flag_insentif_realisasi_kosong"].sum()),
              "Dibiarkan NULL + flag (0 dan NULL beda arti)")

    return df


# ---------------------------------------------------------------- validasi silang
def validasi_referential_integrity(hasil):
    kode_cabang = set(hasil["Data Cabang"]["Branch Code"])
    for sheet, kolom in [("Data Karyawan", "Branch Code"), ("Reason Out", "kode_cabang"),
                         ("Data Kompensasi", "kode_cabang"), ("Performance Cabang", "kode_cabang"),
                         ("Engagement Survey", "kode_cabang")]:
        yatim = set(hasil[sheet][kolom].dropna()) - kode_cabang
        assert not yatim, f"{sheet}.{kolom} punya kode cabang tak dikenal: {yatim}"

    nik = set(hasil["Data Karyawan"]["NIK"])
    for sheet, kolom in [("Performance Data", "nik"), ("Data Kompensasi", "nik")]:
        yatim = set(hasil[sheet][kolom]) - nik
        assert not yatim, f"{sheet}.{kolom} punya NIK tak dikenal: {len(yatim)} NIK"
    print("[OK] Referential integrity antar sheet lolos")


# ---------------------------------------------------------------- output
def tulis_excel(hasil, path):
    urutan = ["Log Pembersihan", "Data Cabang", "Data Karyawan", "Reason Out",
              "Performance Data", "Data Kompensasi", "Performance Cabang", "Engagement Survey"]

    with pd.ExcelWriter(path, engine="openpyxl",
                        datetime_format="yyyy-mm-dd", date_format="yyyy-mm-dd") as w:
        for s in urutan:
            hasil[s].to_excel(w, sheet_name=s, index=False)

    wb = load_workbook(path)
    fill = PatternFill("solid", fgColor="D9D9D9")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 30
        for c in ws[1]:
            c.font = Font(name="Arial", size=10, bold=True)
            c.fill = fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        lebar = {}
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400)):
            for c in row:
                if c.row > 1:
                    c.font = Font(name="Arial", size=10)
                teks = str(c.value) if c.value is not None else ""
                lebar[c.column] = min(max(lebar.get(c.column, 10), len(teks) + 2), 48)
        for col, wd in lebar.items():
            ws.column_dimensions[get_column_letter(col)].width = wd
    wb.save(path)


# ---------------------------------------------------------------- main
def main():
    xls = pd.ExcelFile(SRC)
    raw = {s: xls.parse(s) for s in xls.sheet_names}
    log = LogPembersihan()

    hasil = {}
    hasil["Data Cabang"] = bersihkan_data_cabang(raw["Data Cabang"], log)
    hasil["Data Karyawan"] = bersihkan_data_karyawan(raw["Data Karyawan"],
                                                     raw["Data Kompensasi"], log)
    hasil["Reason Out"] = bersihkan_reason_out(raw["Reason Out"],
                                               set(hasil["Data Karyawan"]["NIK"]), log)
    hasil["Performance Data"] = bersihkan_performance_data(raw["Performance Data"], log)
    hasil["Data Kompensasi"] = bersihkan_data_kompensasi(raw["Data Kompensasi"], log)

    # Dua sheet ini sudah bersih, dilewatkan apa adanya
    hasil["Performance Cabang"] = raw["Performance Cabang"].copy()
    hasil["Engagement Survey"] = raw["Engagement Survey"].copy()
    log.catat("Performance Cabang", "(semua)", "Tidak ditemukan masalah", 0, "-")
    log.catat("Engagement Survey", "(semua)",
              "1 cabang tidak disurvei di 2024 (39 dari 40 cabang)", 1,
              "Dibiarkan (memang tidak ada data)")

    validasi_referential_integrity(hasil)
    hasil["Log Pembersihan"] = log.to_frame()

    tulis_excel(hasil, DST)

    print(log.to_frame().to_string(index=False))
    print(f"\n[OK] Tersimpan: {DST}")
    for s, d in hasil.items():
        print(f"  {s:<20} {d.shape[0]:>6} baris x {d.shape[1]} kolom")


if __name__ == "__main__":
    main()